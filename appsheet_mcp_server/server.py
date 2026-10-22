"""
AppSheet MCP Server for Antigravity AI Agent
Provides full bidirectional integration with Google AppSheet:
- Scaffolding new AppSheet databases (Multi-table Excel with Keys, Refs, Enums)
- Querying data (Find with selector expressions)
- Mutating data (Add, Edit, Delete)
- Invoking custom AppSheet Actions
- Inspecting incoming Webhook events from AppSheet Automation Bots
"""

import os
import sys
import json
import requests
from typing import Any, Dict, List, Optional
from pathlib import Path
from dotenv import load_dotenv

# Load .env if present
env_path = Path(__file__).parent / ".env"
if env_path.exists():
    load_dotenv(dotenv_path=env_path)
else:
    load_dotenv()

from mcp.server.mcpserver import MCPServer

# Initialize MCP Server
mcp = MCPServer("appsheet-agent-mcp")

# File path for incoming webhook events log
EVENTS_LOG_FILE = Path(__file__).parent / "webhook_events.jsonl"


def _get_credentials(app_id: str = "", api_key: str = "") -> tuple[str, str]:
    """Resolve App ID and Access Key from parameters or environment variables."""
    env_file = Path(__file__).resolve().parent / ".env"
    if env_file.exists():
        load_dotenv(dotenv_path=env_file, override=True)

    final_app_id = (app_id or os.getenv("APPSHEET_APP_ID", "")).strip()
    final_api_key = (api_key or os.getenv("APPSHEET_APPLICATION_ACCESS_KEY", "")).strip()
    return final_app_id, final_api_key


def _call_appsheet_api(
    table_name: str,
    action: str,
    rows: Optional[List[Dict[str, Any]]] = None,
    selector: Optional[str] = None,
    app_id: str = "",
    api_key: str = "",
    locale: str = "vi-VN",
    timezone: str = "SE Asia Standard Time"
) -> Dict[str, Any]:
    """Execute standard REST call to AppSheet API v2."""
    final_app_id, final_api_key = _get_credentials(app_id, api_key)

    if not final_app_id or not final_api_key:
        raise ValueError(
            "Thiếu thông tin xác thực AppSheet (APPSHEET_APP_ID hoặc APPSHEET_APPLICATION_ACCESS_KEY). "
            "Vui lòng thiết lập biến môi trường trong .env hoặc truyền trực tiếp vào tham số tool."
        )

    url = f"https://api.appsheet.com/api/v2/apps/{final_app_id}/tables/{table_name}/Action"
    headers = {
        "ApplicationAccessKey": final_api_key,
        "Content-Type": "application/json"
    }

    properties: Dict[str, Any] = {
        "Locale": locale,
        "Timezone": timezone
    }
    if selector:
        properties["Selector"] = selector

    payload: Dict[str, Any] = {
        "Action": action,
        "Properties": properties,
        "Rows": rows or []
    }

    try:
        response = requests.post(url, headers=headers, json=payload, timeout=30)
        response.raise_for_status()
        
        # Check if response is empty
        if not response.text.strip():
            return {"status": "success", "message": "Thao tác thành công nhưng AppSheet không trả về nội dung."}
        
        return response.json()
    except requests.exceptions.HTTPError as e:
        status_code = response.status_code if response else "Unknown"
        error_body = response.text if response else str(e)
        raise RuntimeError(f"Lỗi AppSheet API [{status_code}]: {error_body}")
    except requests.exceptions.RequestException as e:
        raise RuntimeError(f"Lỗi kết nối tới AppSheet: {str(e)}")


@mcp.tool(
    name="appsheet_find_records",
    description="Truy vấn dữ liệu từ bảng trong AppSheet theo điều kiện lọc (Find). Hỗ trợ biểu thức lọc AppSheet Selector."
)
def appsheet_find_records(
    table_name: str,
    selector: str = "",
    app_id: str = "",
    api_key: str = ""
) -> str:
    """
    Truy vấn bản ghi từ bảng AppSheet.
    - table_name: Tên bảng trên AppSheet (vd: 'SanPham', 'DonHang', 'NhanVien')
    - selector: Biểu thức lọc AppSheet (tùy chọn), ví dụ: 'Filter(DonHang, [TrangThai] = "Chờ duyệt")' hoặc '[TrangThai] = "Hoàn thành"'
    - app_id, api_key: Tùy chọn nếu đã cấu hình trong .env
    """
    try:
        result = _call_appsheet_api(
            table_name=table_name,
            action="Find",
            selector=selector if selector else None,
            app_id=app_id,
            api_key=api_key
        )
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_add_rows",
    description="Thêm một hoặc nhiều dòng dữ liệu mới vào bảng AppSheet (Action Add)."
)
def appsheet_add_rows(
    table_name: str,
    rows: List[Dict[str, Any]],
    app_id: str = "",
    api_key: str = ""
) -> str:
    """
    Thêm bản ghi mới vào AppSheet.
    - table_name: Tên bảng đích
    - rows: Danh sách dictionary các trường dữ liệu cần thêm, ví dụ: [{"MaSP": "SP01", "TenSP": "Áo thun", "Gia": 150000}]
    """
    try:
        result = _call_appsheet_api(
            table_name=table_name,
            action="Add",
            rows=rows,
            app_id=app_id,
            api_key=api_key
        )
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_update_rows",
    description="Cập nhật (sửa đổi) một hoặc nhiều dòng dữ liệu trong bảng AppSheet (Action Edit). Bắt buộc phải có giá trị Key ID."
)
def appsheet_update_rows(
    table_name: str,
    rows: List[Dict[str, Any]],
    app_id: str = "",
    api_key: str = ""
) -> str:
    """
    Cập nhật dữ liệu vào bảng AppSheet.
    - table_name: Tên bảng
    - rows: Danh sách các dòng cần cập nhật (mỗi dòng phải chứa trường Key của bảng đó)
    """
    try:
        result = _call_appsheet_api(
            table_name=table_name,
            action="Edit",
            rows=rows,
            app_id=app_id,
            api_key=api_key
        )
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_delete_rows",
    description="Xóa một hoặc nhiều dòng dữ liệu trong bảng AppSheet (Action Delete). Cần truyền trường Key ID."
)
def appsheet_delete_rows(
    table_name: str,
    rows: List[Dict[str, Any]],
    app_id: str = "",
    api_key: str = ""
) -> str:
    """
    Xóa dòng dữ liệu trong AppSheet.
    - table_name: Tên bảng
    - rows: Danh sách chứa trường khóa chính của các dòng cần xóa, ví dụ: [{"ID": "1001"}, {"ID": "1002"}]
    """
    try:
        result = _call_appsheet_api(
            table_name=table_name,
            action="Delete",
            rows=rows,
            app_id=app_id,
            api_key=api_key
        )
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_invoke_action",
    description="Kích hoạt một Custom Action hoặc Grouped Action đã cấu hình trong AppSheet (ví dụ: 'Approve', 'SendAlert', 'ExportPDF')."
)
def appsheet_invoke_action(
    table_name: str,
    action_name: str,
    rows: Optional[List[Dict[str, Any]]] = None,
    app_id: str = "",
    api_key: str = ""
) -> str:
    """
    Kích hoạt action trên AppSheet.
    - table_name: Tên bảng chứa action
    - action_name: Tên chính xác của Action trong AppSheet
    - rows: Danh sách các dòng áp dụng action (nếu action áp dụng trên record cụ thể)
    """
    try:
        result = _call_appsheet_api(
            table_name=table_name,
            action=action_name,
            rows=rows or [],
            app_id=app_id,
            api_key=api_key
        )
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_scaffold_app",
    description="Khởi tạo cấu trúc CSDL hoàn chỉnh chuẩn AppSheet từ ý tưởng (sinh file Excel nhiều sheet, chuẩn Key, Ref, Enum, dữ liệu mẫu, và hướng dẫn tạo app 1-click)."
)
def appsheet_scaffold_app(
    app_name: str,
    business_description: str,
    tables: List[Dict[str, Any]],
    output_dir: str = ""
) -> str:
    """
    Sinh file Excel chuẩn AppSheet từ thiết kế cấu trúc bảng:
    - app_name: Tên ứng dụng (vd: 'QuanLyBaoTriThietBi')
    - business_description: Mô tả nghiệp vụ
    - tables: Danh sách các bảng, mỗi bảng gồm:
        {
          "table_name": "ThietBi",
          "columns": [
             {"name": "MaThietBi", "type": "Text", "is_key": True, "description": "Mã thiết bị"},
             {"name": "TenThietBi", "type": "Text", "is_label": True, "description": "Tên máy"},
             {"name": "LoaiMay", "type": "Enum", "allowed_values": ["MayMay", "MayCat", "MayEp"]},
             {"name": "NhaXuong", "type": "Ref", "ref_table": "NhaXuong"}
          ],
          "sample_rows": [ ... ]
        }
    - output_dir: Thư mục lưu file Excel (mặc định lưu tại thư mục hiện tại)
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        target_dir = Path(output_dir) if output_dir else Path.cwd()
        target_dir.mkdir(parents=True, exist_ok=True)
        
        file_name = f"{app_name.strip().replace(' ', '_')}_Database.xlsx"
        file_path = target_dir / file_name

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style templates
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        key_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        metadata_report = []

        for tbl in tables:
            tbl_name = tbl.get("table_name", "Sheet1")
            columns = tbl.get("columns", [])
            sample_rows = tbl.get("sample_rows", [])

            ws = wb.create_sheet(title=tbl_name)

            col_names = [c["name"] for c in columns]
            ws.append(col_names)

            # Format Header Row
            for col_idx, col_def in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_thin
                
                # Column width
                ws.column_dimensions[cell.column_letter].width = max(len(col_def["name"]) + 6, 16)

            # Append sample rows
            for r_idx, row_data in enumerate(sample_rows, start=2):
                row_vals = []
                for c in columns:
                    val = row_data.get(c["name"], "")
                    row_vals.append(val)
                ws.append(row_vals)
                for col_idx in range(1, len(columns) + 1):
                    cell = ws.cell(row=r_idx, column=col_idx)
                    cell.border = border_thin
                    cell.font = Font(name="Arial", size=10)

            metadata_report.append({
                "table_name": tbl_name,
                "columns_count": len(columns),
                "sample_rows_count": len(sample_rows)
            })

        wb.save(str(file_path))

        # Create quickstart guide markdown
        guide_path = target_dir / f"{app_name}_AppSheet_Setup_Guide.md"
        with open(guide_path, "w", encoding="utf-8") as f:
            f.write(f"# Hướng dẫn tạo App '{app_name}' trên Google AppSheet\n\n")
            f.write(f"**Mô tả:** {business_description}\n\n")
            f.write(f"**File Database:** `{file_path.name}`\n\n")
            f.write("## Các bước kích hoạt AppSheet:\n")
            f.write("1. Tải file Excel này lên **Google Drive** của bạn (hoặc mở bằng Google Sheets).\n")
            f.write("2. Truy cập [AppSheet Create](https://www.appsheet.com/Template/App/CreateFromData).\n")
            f.write(f"3. Chọn nguồn dữ liệu là file `{file_path.name}` vừa tải lên.\n")
            f.write("4. AppSheet sẽ tự động dựng toàn bộ giao diện dựa trên các cột và bảng sau:\n\n")
            for tbl in tables:
                f.write(f"### Bảng: `{tbl.get('table_name')}`\n")
                f.write("| Tên Cột | Kiểu AppSheet | Khóa (Key) | Nhãn (Label) | Ghi chú |\n")
                f.write("| :--- | :--- | :--- | :--- | :--- |\n")
                for c in tbl.get("columns", []):
                    is_k = "✅ Key" if c.get("is_key") else ""
                    is_l = "🏷️ Label" if c.get("is_label") else ""
                    note = c.get("description", "")
                    if c.get("ref_table"):
                        note += f" (Ref tới {c.get('ref_table')})"
                    if c.get("allowed_values"):
                        note += f" (Enum: {', '.join(c.get('allowed_values'))})"
                    f.write(f"| `{c.get('name')}` | {c.get('type', 'Text')} | {is_k} | {is_l} | {note} |\n")
                f.write("\n")

        return json.dumps({
            "status": "success",
            "excel_file": str(file_path),
            "setup_guide": str(guide_path),
            "tables": metadata_report,
            "message": f"Đã sinh cấu trúc CSDL và file Excel thành công cho app '{app_name}'!"
        }, ensure_ascii=False, indent=2)

    except Exception as e:
        return json.dumps({"error": f"Lỗi tạo cấu trúc App: {str(e)}"}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_get_webhook_events",
    description="Kiểm tra các sự kiện (trigger/webhook) gửi từ AppSheet Automation Bot về server."
)
def appsheet_get_webhook_events(limit: int = 10, clear_after_read: bool = False) -> str:
    """
    Lấy danh sách các sự kiện mà người dùng vừa thao tác trên AppSheet (được Bot gửi về webhook).
    - limit: Số lượng sự kiện gần nhất cần lấy (mặc định 10)
    - clear_after_read: Xóa bộ nhớ đệm sau khi đọc nếu cần
    """
    if not EVENTS_LOG_FILE.exists():
        return json.dumps({
            "message": "Chưa có sự kiện nào được ghi nhận. Webhook listener có thể chưa khởi động hoặc chưa có trigger từ AppSheet.",
            "events": []
        }, ensure_ascii=False)

    try:
        events = []
        with open(EVENTS_LOG_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        events.append(json.loads(line))
                    except Exception:
                        pass
        
        recent_events = events[-limit:]
        
        if clear_after_read:
            EVENTS_LOG_FILE.unlink(missing_ok=True)

        return json.dumps({
            "total_captured": len(events),
            "returned_count": len(recent_events),
            "events": recent_events
        }, ensure_ascii=False, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)


@mcp.tool(
    name="appsheet_get_architecture_guide",
    description="Tra cứu hướng dẫn thiết kế chuẩn nghiệp vụ AppSheet (Formulas, Virtual Columns, Ref, Valid_If, Bots)."
)
def appsheet_get_architecture_guide() -> str:
    """Trả về tài liệu quy tắc thiết kế hệ thống AppSheet chuẩn kỹ sư."""
    guide = {
        "Data_Modeling": {
            "Key_Rules": "Mỗi bảng BẮT BUỘC có 1 cột khóa chính (thường là [ID] với Initial Value = UNIQUEID()). Không dùng số thứ tự làm Key.",
            "Ref_Relationships": "Để tạo quan hệ 1-Nhiều (Master-Detail): Cột trong bảng Con mang kiểu Ref trỏ về bảng Cha. Bật 'IsPartOf' nếu muốn form con nhúng trực tiếp trong form cha.",
            "Enum_Columns": "Dùng cho trạng thái quy trình: Ví dụ [TrangThai] gồm ['Mới tạo', 'Chờ duyệt', 'Đã duyệt', 'Từ chối']."
        },
        "Expressions_CheatSheet": {
            "Filter": "FILTER('DonHang', AND([TrangThai] = 'Chờ duyệt', [NguoiTao] = USEREMAIL()))",
            "Sum_Related": "SUM([Related ChiTietPhieus][ThanhTien])",
            "Valid_If": "Đảm bảo tính hợp lệ, vd: [_THIS] > 0 hoặc IN([_THIS], DanhMuc[Ten])",
            "Show_If": "Hiện cột có điều kiện: [LoaiPhieu] = 'Xuất kho'"
        },
        "Automation_Bot_Webhook": {
            "Event": "Data Change (Adds, Updates, Deletes)",
            "Task_Type": "Call a Webhook",
            "HTTP_Verb": "POST",
            "Body_Template": "JSON payload chứa các trường <<[ID]>>, <<[TrangThai]>>, <<USEREMAIL()>>, <<NOW()>>"
        }
    }
    return json.dumps(guide, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    # Run the server via stdio transport (standard for MCP clients like Antigravity)
    mcp.run(transport="stdio")
