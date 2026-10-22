import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
tbl = doc.add_table(rows=2, cols=2)
cell = tbl.cell(0, 0)
tcPr = cell._tc.get_or_add_tcPr()
shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="005BAC"/>')
tcPr.append(shd)

doc.save('test_table.docx')
print('Helper works perfectly!')
